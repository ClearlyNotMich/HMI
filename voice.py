"""Explicit WAV transcription and offline text metrics for the HMI experiment.

Importing this module does not load audio, optional dependencies, or models.
Only the command-line entry point (or run_evaluation) starts model inference.
"""

import argparse
from collections import Counter
from difflib import SequenceMatcher
import math
from pathlib import Path
import re
import time

MODEL_ID = "speechbrain/asr-transformer-transformerlm-librispeech"
REPOSITORY = Path(__file__).resolve().parent
DEFAULT_REFERENCE = (
    REPOSITORY
    / "Speech Recognition Dataset"
    / "speechrecognition Dataset v2"
    / "speechrecognition dataset details.xlsx"
)
DEFAULT_MODEL_CACHE = REPOSITORY / "pretrained_models" / MODEL_ID.split("/")[-1]


def levenshtein_distance(s1, s2):
    if len(s1) < len(s2):
        return levenshtein_distance(s2, s1)

    if len(s2) == 0:
        return len(s1)

    previous_row = range(len(s2) + 1)
    for i, c1 in enumerate(s1):
        current_row = [i + 1]
        for j, c2 in enumerate(s2):
            insertions = previous_row[j + 1] + 1
            deletions = current_row[j] + 1
            substitutions = previous_row[j] + (c1 != c2)
            current_row.append(min(insertions, deletions, substitutions))
        previous_row = current_row

    return previous_row[-1]


def calculate_wer(ref_words, hyp_words):
    """Return exact word edit distance divided by reference token count.

    Pass normalized, unpadded word sequences. Two empty sequences have WER 0;
    a nonempty hypothesis with an empty reference is undefined and raises.
    WER can exceed 1 when the hypothesis contains many insertions.
    """
    if not ref_words:
        if not hyp_words:
            return 0.0
        raise ValueError("WER is undefined for an empty reference")
    return levenshtein_distance(ref_words, hyp_words) / len(ref_words)


def calculate_wer2(ref_words, hyp_words):
    """Compatibility name for the same exact WER calculation."""
    return calculate_wer(ref_words, hyp_words)


def preprocess(text):
    """Keep the original lowercase/punctuation normalization, trimming spaces."""
    text = re.sub(r"[^\w\s]", "", text.lower())
    return re.sub(r"\s+", " ", text).strip()


def similar(a, b):
    # Define a similarity threshold
    threshold = 0.6
    return SequenceMatcher(None, a, b).ratio() > threshold


def find_best_alignment(current_word, other_words, look_ahead_limit=3):
    for offset in range(1, min(look_ahead_limit + 1, len(other_words))):
        if similar(current_word, other_words[offset]):
            return offset
    return 0


def align_texts(reference, hypothesis):
    # Normalize reference and hypothesis texts
    reference = preprocess(reference)
    hypothesis = preprocess(hypothesis)

    reference_words = reference.split()
    hypothesis_words = hypothesis.split()

    align_ref, align_hyp = [], []
    i, j = 0, 0

    while i < len(reference_words) and j < len(hypothesis_words):
        if reference_words[i] == hypothesis_words[j]:
            align_ref.append(reference_words[i])
            align_hyp.append(hypothesis_words[j])
            i += 1
            j += 1
        elif similar(reference_words[i], hypothesis_words[j]):
            align_ref.append(reference_words[i])
            align_hyp.append(hypothesis_words[j])
            i += 1
            j += 1
        else:
            ref_look_ahead = find_best_alignment(
                reference_words[i], hypothesis_words[j:], 3
            )
            hyp_look_ahead = find_best_alignment(
                hypothesis_words[j], reference_words[i:], 3
            )

            if ref_look_ahead > 0:
                align_ref.append("")
                align_hyp.append(hypothesis_words[j])
                j += 1
            elif hyp_look_ahead > 0:
                align_ref.append(reference_words[i])
                align_hyp.append("")
                i += 1
            else:
                align_ref.append(reference_words[i])
                align_hyp.append("")
                i += 1

    # Handle remaining words in either list
    while i < len(reference_words):
        align_ref.append(reference_words[i])
        align_hyp.append("")
        i += 1

    while j < len(hypothesis_words):
        align_ref.append("")
        align_hyp.append(hypothesis_words[j])
        j += 1

    # Printing aligned words side by side for visual confirmation
    print("Heuristic alignment diagnostic (not used for WER):")
    for ref, hyp in zip(align_ref, align_hyp):
        print(f"{ref.ljust(15)} - {hyp}")

    return align_ref, align_hyp


def load_wav_files_from_folder(folder_path):
    """Return sorted WAV filenames from one directory, without decoding audio."""
    folder = Path(folder_path).expanduser().resolve(strict=True)
    if not folder.is_dir():
        raise NotADirectoryError(folder)
    files = sorted(
        path.name
        for path in folder.iterdir()
        if path.is_file() and path.suffix.lower() == ".wav"
    )
    if not files:
        raise ValueError(f"No WAV files found in {folder}")
    return files


def _transcriptions_from_rows(rows):
    """Read task rows up to the bundled workbooks' explicit 'type:' footer."""
    rows = iter(rows)
    headers = next(rows, ())
    if headers.count("ID") != 1 or headers.count("Sentence") != 1:
        raise ValueError("Reference workbook must contain ID and Sentence columns")
    id_index, sentence_index = headers.index("ID"), headers.index("Sentence")
    references = {}
    for number, row in enumerate(rows, 2):
        if not any(value is not None for value in row):
            continue
        if len(row) <= max(id_index, sentence_index):
            raise ValueError(f"Incomplete reference row {number}")
        task_id, sentence = row[id_index], row[sentence_index]
        # The bundled workbooks put speaker metadata below this section marker.
        # Other nonnumeric IDs remain errors; blank rows alone do not hide them.
        if isinstance(task_id, str) and task_id.strip().lower() == "type:":
            break
        if isinstance(task_id, bool):
            raise ValueError(f"Invalid task ID in reference row {number}")
        if isinstance(task_id, str) and task_id.strip().isdigit():
            task_id = int(task_id)
        elif isinstance(task_id, (int, float)) and math.isfinite(task_id):
            if task_id != int(task_id):
                raise ValueError(f"Invalid task ID in reference row {number}")
            task_id = int(task_id)
        else:
            raise ValueError(f"Invalid task ID in reference row {number}")
        if task_id <= 0:
            raise ValueError(f"Invalid task ID in reference row {number}")
        if not isinstance(sentence, str) or not preprocess(sentence).split():
            raise ValueError(f"Empty or invalid sentence in reference row {number}")
        if task_id in references:
            raise ValueError(f"Duplicate reference task ID: {task_id}")
        references[task_id] = sentence
    if not references:
        raise ValueError("Reference workbook contains no transcriptions")
    return references


def load_transcriptions_from_excel(excel_path):
    """Read the first worksheet's ID/Sentence columns; propagate input errors."""
    path = Path(excel_path).expanduser().resolve(strict=True)
    if not path.is_file():
        raise ValueError(f"Reference workbook is not a file: {path}")
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return _transcriptions_from_rows(
            workbook.worksheets[0].iter_rows(values_only=True)
        )
    finally:
        workbook.close()


def select_transcription_for_audio_file(audio_file, transcriptions):
    """Match the original task<number> filename convention to a reference."""
    match = re.search(r"task(\d+)", Path(audio_file).name, flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"WAV filename has no task number: {audio_file}")
    task_id = int(match.group(1))
    if task_id not in transcriptions:
        raise ValueError(f"No reference for task {task_id}: {audio_file}")
    reference = transcriptions[task_id]
    if not isinstance(reference, str) or not preprocess(reference).split():
        raise ValueError(f"Empty reference for task {task_id}")
    return reference


def calculate_recall_precision_fscore(hypothesis, reference):
    """Return aggregate exact-match scores for a supplied heuristic alignment.

    Blank alignment cells are excluded. These are diagnostics of that alignment,
    not macro-averaged scores or the edit-distance alignment used by WER.
    """
    true_positives = sum(
        bool(hyp) and hyp == ref for hyp, ref in zip(hypothesis, reference)
    )
    total_predicted = sum(bool(word) for word in hypothesis)
    total_actual = sum(bool(word) for word in reference)
    precision = true_positives / total_predicted if total_predicted else 0.0
    recall = true_positives / total_actual if total_actual else 0.0
    f_score = (
        2 * precision * recall / (precision + recall) if precision + recall else 0.0
    )
    return recall, precision, f_score


def calculate_for_word(target_words, hypothesis, reference):
    """Print the original target-word diagnostics using nonblank exact matches."""
    hyp_count = Counter(hypothesis)
    ref_count = Counter(reference)
    for word in target_words:
        if not word:
            continue
        tp = sum(hyp == ref == word for hyp, ref in zip(hypothesis, reference))
        recall = tp / ref_count[word] if ref_count[word] else 0.0
        precision = tp / hyp_count[word] if hyp_count[word] else 0.0
        f_score = (
            2 * precision * recall / (precision + recall) if precision + recall else 0.0
        )
        print(
            f"Heuristic aligned word {word!r}: recall={recall}, precision={precision}, F-score={f_score}"
        )


def output(reference, hypothesis):
    """Print exact text metrics separately from the original fuzzy diagnostic."""
    normalized_ref = preprocess(reference)
    normalized_hyp = preprocess(hypothesis)
    wer = calculate_wer(normalized_ref.split(), normalized_hyp.split())
    character_distance = levenshtein_distance(normalized_ref, normalized_hyp)
    print("-----------------")
    print("Reference:", reference)
    print("Hypothesis:", hypothesis)
    print("Word Error Rate (WER):", wer)
    print("Word accuracy (1 - WER; can be negative):", 1 - wer)
    print("Normalized character Levenshtein distance:", character_distance)
    aligned_ref, aligned_hyp = align_texts(reference, hypothesis)
    recall, precision, f_score = calculate_recall_precision_fscore(
        aligned_hyp, aligned_ref
    )
    calculate_for_word(["our", "the"], aligned_hyp, aligned_ref)
    print("Heuristic alignment aggregate recall:", recall)
    print("Heuristic alignment aggregate precision:", precision)
    print("Heuristic alignment aggregate F-score:", f_score)
    return {
        "wer": wer,
        "word_accuracy": 1 - wer,
        "character_distance": character_distance,
    }


def get_audio_duration(audio_path):
    """Validate a WAV header and its duration without loading the waveform."""
    from soundfile import info

    try:
        details = info(str(audio_path))
    except (RuntimeError, OSError) as error:
        raise ValueError(f"Cannot read WAV header: {audio_path}") from error
    if details.format not in {"WAV", "WAVEX", "RF64"}:
        raise ValueError(f"Input is not a WAV audio file: {audio_path}")
    duration = float(details.duration)
    if not math.isfinite(duration) or duration <= 0:
        raise ValueError(f"WAV has no positive finite duration: {audio_path}")
    return duration


def create_asr_model(device, model_cache):
    """Explicitly load the original SpeechBrain model, downloading if needed."""
    if device == "cuda":
        import torch

        if not torch.cuda.is_available():
            raise ValueError("CUDA is unavailable; use --device cpu")
    from speechbrain.inference.ASR import EncoderDecoderASR

    return EncoderDecoderASR.from_hparams(
        source=MODEL_ID,
        savedir=str(model_cache),
        run_opts={"device": device},
    )


def run_evaluation(
    audio_dir,
    reference_path=DEFAULT_REFERENCE,
    *,
    device="cpu",
    model_cache=DEFAULT_MODEL_CACHE,
):
    """Validate every local input, then load one model and evaluate sorted WAVs.

    Only direct children of audio_dir are evaluated. Model-loading time is
    excluded from each transcription's real-time factor, as in the original.
    """
    if device not in {"cpu", "cuda"}:
        raise ValueError("device must be cpu or cuda")
    audio_dir = Path(audio_dir).expanduser().resolve(strict=True)
    filenames = load_wav_files_from_folder(audio_dir)
    references = load_transcriptions_from_excel(reference_path)
    inputs = [
        (audio_dir / name, select_transcription_for_audio_file(name, references))
        for name in filenames
    ]
    model_cache = Path(model_cache).expanduser().resolve()
    if model_cache.exists() and not model_cache.is_dir():
        raise ValueError(f"Model cache is not a directory: {model_cache}")
    # Validate all headers before any model download or inference can begin.
    prepared = [
        (path, reference, get_audio_duration(path)) for path, reference in inputs
    ]
    model = create_asr_model(device, model_cache)
    results = []
    for path, reference, duration in prepared:
        start = time.perf_counter()
        hypothesis = model.transcribe_file(str(path))
        elapsed = time.perf_counter() - start
        print("-----------------")
        print("Audio file:", path.name)
        print("Transcription time (seconds):", elapsed)
        print("Audio duration (seconds):", duration)
        print("Real-time factor (RTF):", elapsed / duration)
        metrics = output(reference, hypothesis)
        results.append(
            {
                "audio_file": path.name,
                "hypothesis": hypothesis,
                "transcription_time": elapsed,
                "audio_duration": duration,
                "rtf": elapsed / duration,
                **metrics,
            }
        )
    return results


def main(argv=None):
    """Run only when explicitly requested; --help needs no optional packages."""
    parser = argparse.ArgumentParser(
        description="Evaluate local WAV files with SpeechBrain (may download model weights).",
    )
    parser.add_argument(
        "--audio-dir",
        required=True,
        type=Path,
        help="directory of task<number>*.wav files",
    )
    parser.add_argument(
        "--reference",
        type=Path,
        default=DEFAULT_REFERENCE,
        help="Excel workbook with ID and Sentence columns",
    )
    parser.add_argument("--device", choices=("cpu", "cuda"), default="cpu")
    parser.add_argument(
        "--model-cache",
        type=Path,
        default=DEFAULT_MODEL_CACHE,
        help="directory for downloaded model files",
    )
    args = parser.parse_args(argv)
    try:
        run_evaluation(
            args.audio_dir,
            args.reference,
            device=args.device,
            model_cache=args.model_cache,
        )
    except (OSError, ValueError) as error:
        parser.error(str(error))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
